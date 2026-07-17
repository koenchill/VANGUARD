"""Convert a Markdown file to a simple Word .docx (headings, tables, lists, code)."""

from __future__ import annotations

import argparse
import re
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

BOLD_OR_CODE = re.compile(r"(\*\*[^*]+\*\*|`[^`]+`)")
LIST_NUM = re.compile(r"^\d+\.\s+")
SEP_CELL = re.compile(r"^:?-{3,}:?$")
IMAGE_LINE = re.compile(r"^!\[([^\]]*)\]\(([^)]+)\)\s*$")
HTML_COMMENT = re.compile(r"^\s*<!--")


def add_runs(paragraph, text: str) -> None:
    for part in BOLD_OR_CODE.split(text):
        if not part:
            continue
        if part.startswith("**") and part.endswith("**"):
            run = paragraph.add_run(part[2:-2])
            run.bold = True
        elif part.startswith("`") and part.endswith("`"):
            run = paragraph.add_run(part[1:-1])
            run.font.name = "Consolas"
            run.font.size = Pt(10)
        else:
            paragraph.add_run(part)


def shade_header_cell(cell) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), "D9E2F3")
    shd.set(qn("w:val"), "clear")
    tc_pr.append(shd)


def add_table(doc: Document, rows: list[list[str]]) -> None:
    if not rows:
        return
    cols = max(len(r) for r in rows)
    table = doc.add_table(rows=len(rows), cols=cols)
    table.style = "Table Grid"
    for i, row in enumerate(rows):
        for j in range(cols):
            cell = table.cell(i, j)
            cell.text = ""
            p = cell.paragraphs[0]
            add_runs(p, row[j] if j < len(row) else "")
            for run in p.runs:
                run.font.size = Pt(10)
                if i == 0:
                    run.bold = True
            if i == 0:
                shade_header_cell(cell)
    doc.add_paragraph()


def parse_table_lines(table_buf: list[str]) -> list[list[str]]:
    rows: list[list[str]] = []
    for tl in table_buf:
        cells = [c.strip() for c in tl.strip().strip("|").split("|")]
        if cells and all(SEP_CELL.match(c.replace(" ", "")) for c in cells):
            continue
        if cells and all(set(c) <= set("-: ") for c in cells):
            continue
        rows.append(cells)
    return rows


def add_image(doc: Document, src_dir: Path, rel_path: str, alt: str) -> None:
    path = Path(rel_path)
    if not path.is_absolute():
        path = (src_dir / path).resolve()
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    if path.is_file():
        run = p.add_run()
        run.add_picture(str(path), width=Inches(6.0))
        cap = doc.add_paragraph()
        cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        cap_run = cap.add_run(alt or path.name)
        cap_run.italic = True
        cap_run.font.size = Pt(9)
    else:
        run = p.add_run(f"[Figure file not found: {rel_path}]")
        run.italic = True
        run.font.color.rgb = RGBColor(0xC0, 0x00, 0x00)


def convert(src: Path, out: Path) -> None:
    text = src.read_text(encoding="utf-8")
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    src_dir = src.parent

    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")

    for level in range(1, 4):
        hs = doc.styles[f"Heading {level}"]
        hs.font.color.rgb = RGBColor(0x1F, 0x1F, 0x1F)
        hs.font.name = "Calibri"

    lines = text.splitlines()
    i = 0
    in_code = False
    in_html_comment = False
    code_buf: list[str] = []
    table_buf: list[str] = []

    def flush_table() -> None:
        nonlocal table_buf
        if table_buf:
            add_table(doc, parse_table_lines(table_buf))
            table_buf = []

    while i < len(lines):
        line = lines[i]

        if in_html_comment:
            if "-->" in line:
                in_html_comment = False
            i += 1
            continue

        if HTML_COMMENT.match(line):
            if "-->" not in line:
                in_html_comment = True
            i += 1
            continue

        if line.strip().startswith("```"):
            if in_code:
                p = doc.add_paragraph()
                run = p.add_run("\n".join(code_buf))
                run.font.name = "Consolas"
                run.font.size = Pt(9)
                p.paragraph_format.left_indent = Inches(0.25)
                code_buf = []
                in_code = False
            else:
                flush_table()
                in_code = True
            i += 1
            continue

        if in_code:
            code_buf.append(line)
            i += 1
            continue

        if line.strip().startswith("|"):
            table_buf.append(line)
            i += 1
            continue

        flush_table()

        if line.strip() == "---":
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(6)
            p.paragraph_format.space_after = Pt(6)
            i += 1
            continue

        if line.startswith("# "):
            title = line[2:].strip()
            # Page break before annexes for stakeholder print/PDF layout
            if title.upper().startswith("ANNEX ") and len(doc.paragraphs) > 3:
                doc.add_page_break()
            p = doc.add_heading(title, level=0)
            if not title.upper().startswith("ANNEX "):
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            i += 1
            continue

        if line.startswith("## "):
            heading = line[3:].strip()
            if heading.lower() in {
                "document control",
                "cover decision banner",
                "table of contents",
            }:
                doc.add_heading(heading, level=1)
            else:
                doc.add_heading(heading, level=1)
            i += 1
            continue

        if line.startswith("### "):
            doc.add_heading(line[4:].strip(), level=2)
            i += 1
            continue

        if line.startswith("#### "):
            doc.add_heading(line[5:].strip(), level=3)
            i += 1
            continue

        stripped = line.strip()
        img = IMAGE_LINE.match(stripped)
        if img:
            add_image(doc, src_dir, img.group(2), img.group(1))
            i += 1
            continue

        if stripped.startswith("[FIGURE ") and stripped.endswith("]"):
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = p.add_run(stripped)
            run.italic = True
            run.font.color.rgb = RGBColor(0x66, 0x66, 0x66)
            i += 1
            continue

        if LIST_NUM.match(stripped):
            p = doc.add_paragraph(style="List Number")
            add_runs(p, LIST_NUM.sub("", stripped))
            i += 1
            continue

        if stripped.startswith("- ") or stripped.startswith("* "):
            p = doc.add_paragraph(style="List Bullet")
            add_runs(p, stripped[2:])
            i += 1
            continue

        if (
            stripped.startswith("*")
            and stripped.endswith("*")
            and not stripped.startswith("**")
        ):
            p = doc.add_paragraph()
            run = p.add_run(stripped.strip("*").strip())
            run.italic = True
            i += 1
            continue

        if not stripped:
            i += 1
            continue

        p = doc.add_paragraph()
        add_runs(p, stripped)
        i += 1

    flush_table()
    out.parent.mkdir(parents=True, exist_ok=True)
    doc.save(out)


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("src", type=Path)
    parser.add_argument("-o", "--output", type=Path, default=None)
    args = parser.parse_args()
    out = args.output or args.src.with_suffix(".docx")
    convert(args.src, out)
    print(f"Wrote {out.resolve()} ({out.stat().st_size} bytes)")


if __name__ == "__main__":
    main()
